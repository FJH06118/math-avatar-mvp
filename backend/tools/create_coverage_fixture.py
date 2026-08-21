"""Create synthetic PPTX files for stage 11B page-count coverage tests."""

from __future__ import annotations

import argparse
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt


def create_fixture(
    output: Path,
    slide_count: int,
    *,
    layout: str = "wide",
    blank_slide: int = 0,
) -> None:
    if not 0 <= slide_count <= 100:
        raise ValueError("slide_count must be between 0 and 100")
    if layout not in {"wide", "standard", "portrait"}:
        raise ValueError("layout must be wide, standard, or portrait")
    if blank_slide < 0 or blank_slide > slide_count:
        raise ValueError("blank_slide must be zero or a valid slide number")
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    if layout == "wide":
        presentation.slide_width = Inches(13.333333)
        presentation.slide_height = Inches(7.5)
    elif layout == "standard":
        presentation.slide_width = Inches(10)
        presentation.slide_height = Inches(7.5)
    else:
        presentation.slide_width = Inches(7.5)
        presentation.slide_height = Inches(10)
    formula_image = output.with_suffix(".formula.png")
    image = Image.new("RGB", (640, 180), "white")
    ImageDraw.Draw(image).text((24, 68), "f'(x) = lim h->0 [f(x+h)-f(x)]/h", fill="black")
    image.save(formula_image)
    try:
        for index in range(1, slide_count + 1):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            if index == blank_slide:
                continue
            content_width = 11.8 if layout == "wide" else 8.6 if layout == "standard" else 6.2
            title = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(content_width), Inches(0.7))
            run = title.text_frame.paragraphs[0].add_run()
            run.text = f"覆盖测试第 {index} 页"
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(28)
            body_width = 11.2 if layout == "wide" else 8.2 if layout == "standard" else 5.8
            body = slide.shapes.add_textbox(Inches(0.9), Inches(1.7), Inches(body_width), Inches(2.0))
            body.text_frame.text = f"稳定页码 {index}/{slide_count}。公式候选：f(x)=x^{index % 5 + 1}。"
            if index == 1:
                table_width = 5.8 if layout == "portrait" else 7.2
                table = slide.shapes.add_table(
                    2,
                    2,
                    Inches(0.9),
                    Inches(3.4),
                    Inches(table_width),
                    Inches(1.1),
                ).table
                table.cell(0, 0).text = "结构"
                table.cell(0, 1).text = "兼容性"
                table.cell(1, 0).text = "表格"
                table.cell(1, 1).text = "可提取"
            if index == min(2, slide_count):
                picture_left = 0.75 if layout == "portrait" else 1.4
                picture_width = 6 if layout == "portrait" else 7.2
                slide.shapes.add_picture(
                    str(formula_image),
                    Inches(picture_left),
                    Inches(5.0),
                    width=Inches(picture_width),
                )
        presentation.save(output)
    finally:
        formula_image.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True)
    parser.add_argument("--slides", type=int, required=True)
    parser.add_argument(
        "--layout",
        choices=("wide", "standard", "portrait"),
        default="wide",
    )
    parser.add_argument("--blank-slide", type=int, default=0)
    args = parser.parse_args()
    create_fixture(
        Path(args.output).resolve(),
        args.slides,
        layout=args.layout,
        blank_slide=args.blank_slide,
    )


if __name__ == "__main__":
    main()
